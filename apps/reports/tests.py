import csv
import io
from datetime import datetime, time, timedelta
from decimal import Decimal

from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.roles import Role
from apps.core.models import AuditLog
from apps.core.testing import HotelTestCase
from apps.helpdesk import services as ticket_services
from apps.helpdesk.models import Ticket, TicketCategory
from apps.projects import services as project_services
from apps.projects.models import Project
from apps.tasks import services as task_services
from apps.tasks.models import Task, TaskCategory

from . import analytics, builders, exporters
from .context import ReportContext
from .structure import Metric, Report, Section

S = Ticket.Status


class ReportTestCase(HotelTestCase):
    def setUp(self):
        super().setUp()
        self.emp = self.make_user("rose")
        self.emp_other = self.make_user("chris", department=self.other_dept)
        self.mgr = self.make_user("fatma", Role.MANAGER)
        self.it = self.make_user("ivan", Role.IT_OFFICER, first_name="Ivan", last_name="Zed")
        self.it2 = self.make_user("anna", Role.IT_OFFICER, first_name="Anna", last_name="Abel")
        self.admin = self.make_user("adm", Role.ADMIN)
        self.cat = TicketCategory.objects.get(name="Printer / scanner", parent=None)
        self.task_cat = TaskCategory.objects.get(name="User support")
        self.today = timezone.localdate()

    def at(self, hh, day=None):
        return timezone.make_aware(datetime.combine(day or self.today, time(hh)))

    def ticket(self, user=None, **kw):
        kw.setdefault("subject", "Printer not printing")
        kw.setdefault("description", "Shows error E-05")
        kw.setdefault("category", self.cat)
        return ticket_services.create_ticket(Ticket(**kw), user or self.emp)

    def resolve(self, ticket, by=None):
        by = by or self.it
        ticket_services.change_status(ticket, S.IN_PROGRESS, by)
        return ticket_services.change_status(ticket, S.RESOLVED, by, resolution_details="Cleared paper jam")

    def task(self, **kw):
        kw.setdefault("title", "Check backups")
        kw.setdefault("category", self.task_cat)
        kw.setdefault("assigned_to", self.it)
        kw.setdefault("planned_date", self.today)
        return task_services.create_task(Task(**kw), self.it)

    def ctx(self, kind="custom", user=None, **kw):
        kw.setdefault("start", self.today - timedelta(days=6))
        kw.setdefault("end", self.today)
        return ReportContext(kind=kind, user=user or self.admin, **kw)

    def metric(self, report, label):
        for s in report.sections:
            for m in s.metrics:
                if m.label == label:
                    return m
        raise AssertionError(f"metric {label!r} not found")

    def section(self, report, title):
        return next((s for s in report.sections if s.title == title), None)


class ReportFiguresTests(ReportTestCase):
    def test_ticket_counts_and_confirmed_resolution_is_verified(self):
        t1, t2 = self.ticket(), self.ticket(subject="Wi-Fi down")
        self.ticket(user=self.emp_other, subject="Till offline")
        self.resolve(t1)
        self.resolve(t2)
        ticket_services.confirm_resolution(t1, self.emp, rating=5)
        report = builders.build_weekly(self.ctx("weekly", start=self.today - timedelta(days=self.today.weekday())))
        self.assertEqual(self.metric(report, "Requests received").value, 3)
        self.assertEqual(self.metric(report, "Requests resolved").value, 2)
        confirmed = self.metric(report, "Resolved and confirmed by requester")
        self.assertEqual((confirmed.value, confirmed.basis), (1, "verified"))
        self.assertEqual(self.metric(report, "Time recorded").basis, "reported")

    def test_linked_task_time_is_counted_once(self):
        t = self.ticket()
        task = self.task(related_ticket=t)
        task_services.log_time(task, self.it, started_at=self.at(9), minutes=40)
        ticket_services.log_time(t, self.it, started_at=self.at(10), minutes=20)
        self.assertEqual(self.metric(builders.build_daily(self.ctx("daily", start=self.today, end=self.today)),
                                     "Time recorded").value, "1 h")
        data = analytics.build(self.ctx())
        self.assertEqual(data["time"]["total_minutes"], 60)

    def test_planned_late_and_unplanned_figures(self):
        done = self.task()
        task_services.change_status(done, Task.Status.COMPLETED, self.it, completion_notes="All good")
        late = self.task(planned_date=self.today - timedelta(days=2), title="Patch server")
        task_services.change_status(late, Task.Status.COMPLETED, self.it, completion_notes="Done late")
        self.task(title="Open job")
        tf = analytics.build(self.ctx())["tf"]
        self.assertEqual((tf["planned"], tf["planned_done"], tf["late"]), (3, 2, 1))

    def test_savings_only_listed_when_documented(self):
        project_services.create_project(Project(
            name="Printer consolidation", description="x", business_objective="y", owner=self.it,
            status=Project.Status.COMPLETED, completed_on=self.today, documented_savings=Decimal("1200"),
            savings_evidence="Toner invoices before/after", currency="USD"), self.it)
        project_services.create_project(Project(
            name="Wi-Fi refresh", description="x", business_objective="y", owner=self.it,
            status=Project.Status.COMPLETED, completed_on=self.today, currency="USD"), self.it)
        report = builders.build_monthly(self.ctx("monthly", start=self.today.replace(day=1)))
        savings = self.section(report, "Documented cost savings")
        self.assertEqual(len(savings.rows), 1)
        self.assertIn("Printer consolidation", savings.rows[0][0])
        self.assertEqual(savings.rows[0][3], "Reported by IT")

    def test_every_kind_builds(self):
        self.resolve(self.ticket())
        for kind, build in builders.BUILDERS.items():
            start, end = builders.default_period(kind)
            report = build(self.ctx(kind, start=start, end=end))
            self.assertTrue(report.sections, kind)
            self.assertTrue(report.definitions, kind)


class ScopeTests(ReportTestCase):
    def test_manager_is_forced_to_own_department(self):
        ctx = self.ctx(user=self.mgr, department_id=self.other_dept.pk)
        self.assertEqual(ctx.department_id, self.dept.pk)
        self.ticket()
        self.ticket(user=self.emp_other)
        report = builders.build_custom(ctx)
        self.assertEqual(self.metric(report, "Requests received").value, 1)

    def test_manager_does_not_see_individual_staff_sections(self):
        report = builders.build_daily(self.ctx("daily", user=self.mgr, start=self.today, end=self.today))
        self.assertFalse(any(s.it_only for s in report.sections))
        self.assertIsNone(self.section(report, "Recorded time per IT staff member"))
        self.assertEqual(analytics.build(self.ctx(user=self.mgr))["officers"], [])

    def test_employee_cannot_open_reports(self):
        self.login(self.emp)
        self.assertEqual(self.client.get(reverse("reports:centre"), {"kind": "daily", "format": "csv"}).status_code, 403)
        self.assertEqual(self.client.get(reverse("reports:analytics")).status_code, 403)

    def test_manager_department_in_url_is_ignored(self):
        self.ticket(user=self.emp_other, subject="Other department issue")
        self.login(self.mgr)
        resp = self.client.get(reverse("reports:centre"), {"kind": "custom", "department": self.other_dept.pk})
        self.assertEqual(resp.status_code, 200)
        self.assertNotContains(resp, "Other department issue")


class ExportTests(ReportTestCase):
    def test_all_formats_download_and_are_audited(self):
        self.resolve(self.ticket())
        self.login(self.it)
        for kind in builders.BUILDERS:
            resp = self.client.get(reverse("reports:centre"), {"kind": kind, "format": "html"})
            self.assertContains(resp, "Definitions")
            pdf = self.client.get(reverse("reports:centre"), {"kind": kind, "format": "pdf"})
            self.assertTrue(pdf.content.startswith(b"%PDF"), kind)
            self.assertIn("attachment", pdf["Content-Disposition"])
            self.assertEqual(pdf["Cache-Control"], "private, no-store")
            wb = load_workbook(io.BytesIO(self.client.get(reverse("reports:centre"),
                                                          {"kind": kind, "format": "xlsx"}).content))
            self.assertIn("Summary", wb.sheetnames)
            self.assertIn("Definitions", wb.sheetnames)
            text = self.client.get(reverse("reports:centre"), {"kind": kind, "format": "csv"}).content.decode("utf-8-sig")
            self.assertTrue(list(csv.reader(io.StringIO(text))))
        self.assertEqual(AuditLog.objects.filter(action=AuditLog.Action.REPORT).count(), 4 * len(builders.BUILDERS))

    def test_formula_injection_is_neutralised(self):
        report = Report(kind="custom", title="T", period_label="P", filters="", generated_at=timezone.now(),
                        generated_by="x", hotel_name="H")
        report.add(Section("Rows", "table", columns=["A"], rows=[["=HYPERLINK(\"http://evil\")"], ["+1"], ["ok"]]))
        report.add(Section("M", "metrics", metrics=[Metric("@cmd", 1)]))
        text = exporters.to_csv(report).decode("utf-8-sig")
        self.assertIn("'=HYPERLINK", text)
        self.assertIn("'+1", text)
        self.assertIn("'@cmd", text)
        wb = load_workbook(io.BytesIO(exporters.to_xlsx(report)))
        values = [c.value for ws in wb for row in ws.iter_rows() for c in row if c.value]
        self.assertFalse(any(isinstance(v, str) and v.startswith("=") for v in values))
        self.assertTrue(exporters.to_pdf(report).startswith(b"%PDF"))

    def test_invalid_form_shows_centre(self):
        self.login(self.it)
        resp = self.client.get(reverse("reports:centre"), {"kind": "bogus"})
        self.assertEqual(resp.status_code, 200)
        self.assertTemplateUsed(resp, "reports/centre.html")

    def test_period_is_capped(self):
        self.login(self.it)
        resp = self.client.get(reverse("reports:centre"), {"kind": "weekly", "start": "2020-01-01", "end": "2026-01-01"})
        self.assertEqual((resp.context["ctx"].end - resp.context["ctx"].start).days, 30)


class AnalyticsTests(ReportTestCase):
    def test_page_renders_with_fairness_note_and_alphabetical_staff(self):
        self.resolve(self.ticket(), by=self.it)
        self.login(self.admin)
        resp = self.client.get(reverse("reports:analytics"))
        self.assertContains(resp, "Listed alphabetically, not ranked")
        names = [row[0] for row in resp.context["d"]["officers"]]
        self.assertEqual(names, sorted(names, key=str.lower))
        self.assertLess(names.index("Anna Abel"), names.index("Ivan Zed"))

    def test_recurring_incident_rate(self):
        self.ticket(ticket_type="incident")
        self.ticket(ticket_type="incident")
        rate, repeats, incidents = analytics.recurring_incident_rate(self.ctx())
        self.assertEqual((repeats, incidents), (1, 2))
        self.assertEqual(rate, 50)

    def test_insights_report_rising_workload(self):
        trend = [{"received": 5, "resolved": 5, "interruptions": 0, "unplanned_hours": 0, "backlog": 2,
                  "planned": 0, "planned_done": 0} for _ in range(8)]
        for row in trend[-4:]:
            row["received"] = 10
        out = analytics.insights(trend, self.ctx())
        self.assertTrue(any("Requests received rose by 100%" in line for line in out))

    def test_backlog_counts_waiting_separately(self):
        t = self.ticket()
        ticket_services.change_status(t, S.WAITING, self.it, waiting_on="vendor", note="Awaiting part")
        b = analytics.backlog(self.ctx())
        self.assertEqual((b["tickets"], b["waiting"]), (1, 1))

    def test_form_shows_default_period_and_long_periods_are_capped(self):
        self.login(self.it)
        resp = self.client.get(reverse("reports:analytics"))
        self.assertContains(resp, f'value="{(self.today - timedelta(days=89)):%Y-%m-%d}"')
        resp = self.client.get(reverse("reports:analytics"), {"kind": "custom", "start": "2015-01-01",
                                                              "end": f"{self.today:%Y-%m-%d}"})
        self.assertEqual((resp.context["ctx"].end - resp.context["ctx"].start).days, 399)
