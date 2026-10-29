"""PDF (ReportLab), Excel (openpyxl) and CSV renderers for a ``Report``."""
import csv
import io

from django.utils import timezone

from .structure import BASIS

BRAND = "#14213D"
ACCENT = "#FCA311"


def _cell(value):
    """Plain text for exports; neutralises spreadsheet formula injection."""
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return value
    text = str(value)
    if text[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + text
    return text


def _stamp(report):
    return timezone.localtime(report.generated_at).strftime("%d %b %Y %H:%M")


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------
def to_csv(report):
    buf = io.StringIO()
    buf.write("﻿")
    w = csv.writer(buf)
    w.writerow([report.hotel_name, report.title])
    w.writerow(["Period", report.period_label])
    w.writerow(["Filters", report.filters])
    w.writerow(["Generated", _stamp(report), "by", report.generated_by])
    for s in report.sections:
        w.writerow([])
        w.writerow([s.title.upper()])
        if s.kind == "metrics":
            w.writerow(["Measure", "Value", "Basis"])
            for m in s.metrics:
                w.writerow([_cell(m.label), _cell(m.value), BASIS[m.basis][0]])
        elif s.kind == "table":
            w.writerow(s.columns)
            for row in s.rows or [[s.empty]]:
                w.writerow([_cell(v) for v in row])
        elif s.kind == "text":
            for heading, body in s.paragraphs or [("", s.empty)]:
                w.writerow([_cell(heading), _cell(body)])
        elif s.kind == "chart" and s.chart:
            w.writerow(["", *s.chart["labels"]])
            for ds in s.chart["datasets"]:
                w.writerow([ds["label"], *ds["data"]])
        if s.note:
            w.writerow(["Note", _cell(s.note)])
    w.writerow([])
    w.writerow(["DEFINITIONS"])
    for term, text in report.definitions:
        w.writerow([term, text])
    return buf.getvalue().encode("utf-8")


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def to_xlsx(report):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=BRAND[1:])
    ws.append([report.hotel_name])
    ws.append([report.title])
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"].font = Font(bold=True, size=12)
    ws.append(["Period", report.period_label])
    ws.append(["Filters", report.filters])
    ws.append(["Generated", f"{_stamp(report)} by {report.generated_by}"])
    used = set()

    def sheet_name(title):
        base = "".join(c for c in title if c not in "[]:*?/\\")[:28] or "Section"
        name, i = base, 2
        while name in used:
            name = f"{base[:25]} {i}"
            i += 1
        used.add(name)
        return name

    for s in report.sections:
        if s.kind == "metrics":
            ws.append([])
            ws.append([s.title])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            for m in s.metrics:
                ws.append([_cell(m.label), _cell(m.value), BASIS[m.basis][0]])
            if s.note:
                ws.append(["Note", s.note])
        elif s.kind == "text":
            ws.append([])
            ws.append([s.title])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            for heading, body in s.paragraphs or [("", s.empty)]:
                ws.append([heading, body])
                ws.cell(ws.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        elif s.kind == "table":
            t = wb.create_sheet(sheet_name(s.title))
            t.append([s.title])
            t["A1"].font = Font(bold=True, size=12)
            if s.note:
                t.append([s.note])
            t.append(s.columns)
            header_row = t.max_row
            for c in t[header_row]:
                c.font, c.fill = head, fill
            for row in s.rows:
                t.append([_cell(v) for v in row])
            if not s.rows:
                t.append([s.empty])
            for i, col in enumerate(s.columns, 1):
                width = max([len(str(col))] + [len(str(r[i - 1])) for r in s.rows[:200] if i - 1 < len(r)])
                t.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
            t.freeze_panes = t.cell(header_row + 1, 1)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 18
    d = wb.create_sheet("Definitions")
    d.append(["Term", "Definition"])
    for term, text in report.definitions + [(v[0], v[1]) for v in BASIS.values()]:
        d.append([term, text])
    d.column_dimensions["A"].width = 28
    d.column_dimensions["B"].width = 100
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def to_pdf(report):
    from reportlab.graphics.charts.barcharts import HorizontalBarChart, VerticalBarChart
    from reportlab.graphics.charts.legends import Legend
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from reportlab.pdfgen import canvas as rl_canvas
    from xml.sax.saxutils import escape

    brand, accent = colors.HexColor(BRAND), colors.HexColor(ACCENT)
    palette = [colors.HexColor(c) for c in ("#1f6feb", "#d9480f", "#2b8a3e", "#ae3ec9", "#f59f00", "#0b7285")]
    styles = getSampleStyleSheet()
    base = ParagraphStyle("base", parent=styles["Normal"], fontSize=8.5, leading=11)
    small = ParagraphStyle("small", parent=base, fontSize=7.5, leading=9.5, textColor=colors.HexColor("#555555"))
    h1 = ParagraphStyle("h1", parent=styles["Title"], fontSize=17, leading=21, textColor=brand, alignment=0, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=11.5, leading=14, textColor=brand, spaceBefore=10,
                        spaceAfter=4)
    cell = ParagraphStyle("cell", parent=base, fontSize=7.5, leading=9)
    cell_head = ParagraphStyle("cellh", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    width = A4[0] - 30 * mm

    class NumberedCanvas(rl_canvas.Canvas):
        """Adds 'Page X of Y', the hotel name and a confidentiality line to every page."""

        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._saved = []

        def showPage(self):
            self._saved.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._saved)
            for state in self._saved:
                self.__dict__.update(state)
                self.setFillColor(brand)
                self.rect(0, A4[1] - 12 * mm, A4[0], 12 * mm, stroke=0, fill=1)
                self.setFillColor(accent)
                self.rect(0, A4[1] - 13 * mm, A4[0], 1 * mm, stroke=0, fill=1)
                self.setFillColor(colors.white)
                self.setFont("Helvetica-Bold", 10)
                self.drawString(15 * mm, A4[1] - 8 * mm, f"{report.hotel_name} · IT Department")
                self.setFont("Helvetica", 8)
                self.drawRightString(A4[0] - 15 * mm, A4[1] - 8 * mm, report.title)
                self.setFillColor(colors.HexColor("#666666"))
                self.setFont("Helvetica", 7.5)
                self.drawString(15 * mm, 9 * mm, f"Confidential – internal use. Generated {_stamp(report)}.")
                self.drawRightString(A4[0] - 15 * mm, 9 * mm, f"Page {self._pageNumber} of {total}")
                super().showPage()
            super().save()

    def para(text, style=base):
        return Paragraph(escape(str(text)).replace("\n", "<br/>"), style)

    def table(columns, rows, col_widths=None):
        data = [[para(c, cell_head) for c in columns]] + [[para(v, cell) for v in row] for row in rows]
        t = Table(data, colWidths=col_widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), brand), ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F5F9")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CCCCCC")),
            ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
        return t

    def chart(spec):
        labels = [str(x)[:28] for x in spec["labels"]]
        datasets = spec["datasets"]
        if spec["type"] == "doughnut":
            spec = {"type": "bar", "horizontal": True, "labels": labels,
                    "datasets": [{"label": datasets[0]["label"], "data": datasets[0]["data"]}]}
        horizontal = spec.get("horizontal")
        height = max(120, 18 * len(labels) + 50) if horizontal else 170
        d = Drawing(width, height)
        if spec["type"] == "line":
            c = HorizontalLineChart()
        else:
            c = HorizontalBarChart() if horizontal else VerticalBarChart()
        c.x, c.y = (110 if horizontal else 30), 30
        c.width, c.height = width - c.x - 20, height - 55
        c.data = [[float(v or 0) for v in ds["data"]] for ds in spec["datasets"]]
        c.categoryAxis.categoryNames = labels
        c.categoryAxis.labels.fontSize = 6.5
        if not horizontal and len(labels) > 8:
            c.categoryAxis.labels.angle, c.categoryAxis.labels.boxAnchor = 45, "ne"
        c.valueAxis.valueMin = 0
        c.valueAxis.labels.fontSize = 6.5
        for i in range(len(c.data)):
            if spec["type"] == "line":
                c.lines[i].strokeColor = palette[i % len(palette)]
                c.lines[i].strokeWidth = 1.5
            else:
                c.bars[i].fillColor = palette[i % len(palette)]
                c.bars[i].strokeColor = None
        d.add(c)
        if len(spec["datasets"]) > 1:
            legend = Legend()
            legend.x, legend.y = c.x, height - 8
            legend.fontSize, legend.columnMaximum = 7, 1
            legend.deltax, legend.dx, legend.dy = 70, 6, 6
            legend.alignment = "right"
            legend.colorNamePairs = [(palette[i % len(palette)], ds["label"]) for i, ds in enumerate(spec["datasets"])]
            d.add(legend)
        return d

    story = [para(report.title, h1),
             Paragraph(f"<b>Period:</b> {escape(report.period_label)}", base),
             para(f"Filters: {report.filters}", small),
             para(f"Generated {_stamp(report)} by {report.generated_by}", small), Spacer(1, 6)]
    for s in report.sections:
        block = [para(s.title, h2)]
        if s.kind == "metrics":
            rows = [[m.label, str(m.value), BASIS[m.basis][0]] for m in s.metrics]
            block.append(table(["Measure", "Value", "Basis"], rows, [width * 0.5, width * 0.3, width * 0.2]))
        elif s.kind == "table":
            if s.rows:
                n = len(s.columns)
                block.append(table(s.columns, s.rows, [width / n] * n))
            else:
                block.append(para(s.empty, small))
        elif s.kind == "text":
            for heading, body in s.paragraphs or [("", s.empty)]:
                if heading:
                    block.append(Paragraph(f"<b>{escape(heading)}</b>", base))
                block.append(para(body, base))
                block.append(Spacer(1, 3))
        elif s.kind == "chart" and s.chart:
            block.append(chart(s.chart))
        if s.note:
            block.append(Spacer(1, 2))
            block.append(para(s.note, small))
        if s.kind in ("metrics", "chart", "text") or len(s.rows) < 12:
            story.append(KeepTogether(block))  # short blocks never split across pages
        else:
            story.extend(block)  # long tables flow, repeating their header row
    story.append(para("Definitions", h2))
    defs = report.definitions + [(f"Basis: {label}", text) for _k, label, text in report.bases_used]
    story.append(table(["Term", "Meaning"], defs, [width * 0.25, width * 0.75]))
    right = ParagraphStyle("right", parent=small, alignment=TA_RIGHT)
    story.append(Spacer(1, 6))
    story.append(para("All figures are calculated from records in the IT management system at the time of generation. "
                      "Nothing is estimated or inferred unless labelled “Estimated”.", right))
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=20 * mm,
                            bottomMargin=16 * mm, title=report.title, author=f"{report.hotel_name} IT")
    doc.build(story, canvasmaker=NumberedCanvas)
    return out.getvalue()
