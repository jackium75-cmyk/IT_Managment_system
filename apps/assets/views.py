import csv
from datetime import timedelta

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.paginator import Paginator
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST

from apps.accounts.models import Department, Location
from apps.accounts.roles import IT_STAFF_ROLES, Role, is_it_staff
from apps.core import audit
from apps.core.models import AuditLog
from apps.core.permissions import role_required
from apps.evidence import serving
from apps.tasks.views import _apply_errors

from . import services
from .forms import AssetCategoryForm, AssetForm, AttachmentForm, StatusForm, SupplierForm
from .models import Asset, AssetAttachment, AssetCategory, Supplier

it_required = role_required(*IT_STAFF_ROLES)
PAGE_SIZE = 30
SORTS = {"tag": ("tag",), "-tag": ("-tag",), "name": ("name",), "category": ("category__name", "tag"),
         "location": ("location__name", "tag"), "status": ("status", "tag"), "warranty": ("warranty_expiry",),
         "purchased": ("purchase_date",)}
EXPORT_COLUMNS = [
    ("Asset ID", "tag"), ("Name", "name"), ("Category", "category"), ("Manufacturer", "manufacturer"),
    ("Model", "model"), ("Serial number", "serial_number"), ("Status", "get_status_display"),
    ("Condition", "get_condition_display"), ("Department", "department"), ("Location", "location"),
    ("Assigned to", "assigned_name"), ("Purchase date", "purchase_date"), ("Purchase cost", "purchase_cost"),
    ("Currency", "currency"), ("Supplier", "supplier"), ("Warranty expiry", "warranty_expiry"),
    ("Hostname", "hostname"), ("IP address", "ip_address"), ("Operating system", "operating_system"),
    ("OS version", "os_version"), ("Replacement", "get_replacement_display"),
]


def _filtered(request):
    qs = Asset.objects.select_related("category", "department", "location", "assigned_to", "supplier")
    f = {k: request.GET.get(k, "").strip() for k in
         ("q", "category", "status", "department", "location", "condition", "assigned", "warranty", "attention")}
    if f["q"]:
        q = f["q"]
        qs = qs.filter(Q(tag__icontains=q) | Q(name__icontains=q) | Q(serial_number__icontains=q) | Q(model__icontains=q)
                       | Q(manufacturer__icontains=q) | Q(hostname__icontains=q) | Q(ip_address__icontains=q)
                       | Q(assigned_to__first_name__icontains=q) | Q(assigned_to__last_name__icontains=q))
    if f["category"].isdigit():
        qs = qs.filter(category_id=f["category"])
    if f["status"] == "in_service" or not f["status"]:
        qs = qs.in_service()
    elif f["status"] in Asset.Status.values:
        qs = qs.filter(status=f["status"])
    if f["department"].isdigit():
        qs = qs.filter(department_id=f["department"])
    if f["location"].isdigit():
        qs = qs.filter(location_id=f["location"])
    if f["condition"] in Asset.Condition.values:
        qs = qs.filter(condition=f["condition"])
    if f["assigned"].isdigit():
        qs = qs.filter(assigned_to_id=f["assigned"])
    today = timezone.localdate()
    if f["warranty"] == "expiring":
        qs = qs.filter(warranty_expiry__range=(today, today + timedelta(days=60)))
    elif f["warranty"] == "expired":
        qs = qs.filter(warranty_expiry__lt=today)
    elif f["warranty"] == "valid":
        qs = qs.filter(warranty_expiry__gte=today)
    if f["attention"] == "1":
        qs = qs.filter(pk__in=[a.pk for a, _r in services.needing_attention(qs)])
    sort = request.GET.get("sort", "tag")
    return qs.order_by(*SORTS.get(sort, SORTS["tag"])), f, sort


@it_required
def asset_list(request):
    qs, f, sort = _filtered(request)
    page = Paginator(qs, PAGE_SIZE).get_page(request.GET.get("page"))
    attention = {a.pk: r for a, r in services.needing_attention(Asset.objects.filter(pk__in=[a.pk for a in page]))}
    return render(request, "assets/asset_list.html", {
        "page_obj": page, "f": f, "sort": sort, "attention": attention,
        "categories": AssetCategory.objects.all(), "statuses": Asset.Status.choices, "conditions": Asset.Condition.choices,
        "departments": Department.objects.active(), "locations": Location.objects.active(),
        "export_query": request.GET.urlencode(),
    })


@it_required
def asset_export(request, fmt):
    qs, _f, _s = _filtered(request)
    rows = []
    for a in qs:
        a.assigned_name = (a.assigned_to.get_full_name() or a.assigned_to.username) if a.assigned_to else ""
        row = []
        for _label, attr in EXPORT_COLUMNS:
            value = getattr(a, attr)
            value = value() if callable(value) else value
            row.append("" if value is None else (value if isinstance(value, (int, float)) else str(value)))
        rows.append(row)
    stamp = timezone.localtime().strftime("%Y%m%d-%H%M")
    audit.log(AuditLog.Action.REPORT, message=f"Asset register exported ({fmt}, {len(rows)} rows)")
    headers = [c[0] for c in EXPORT_COLUMNS]
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Asset register"
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([_safe_cell(v) for v in row])
        for i, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, len(header) + 2)
        ws.freeze_panes = "A2"
        resp = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp["Content-Disposition"] = f'attachment; filename="asset-register-{stamp}.xlsx"'
        wb.save(resp)
        return resp
    resp = HttpResponse(content_type="text/csv; charset=utf-8")
    resp["Content-Disposition"] = f'attachment; filename="asset-register-{stamp}.csv"'
    resp.write("﻿")  # so Excel opens UTF-8 correctly
    writer = csv.writer(resp)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([_safe_cell(v) for v in row])
    return resp


def _safe_cell(value):
    """Neutralise spreadsheet formula injection (=, +, -, @ at the start of text)."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


@it_required
def asset_create(request):
    if request.method == "POST":
        form = AssetForm(request.POST)
        if form.is_valid():
            try:
                asset = services.create_asset(form.save(commit=False), request.user)
            except ValidationError as exc:
                _apply_errors(form, exc)
            else:
                messages.success(request, f"Asset registered as {asset.tag}. Print its QR label from the asset page.")
                return redirect(asset)
    else:
        form = AssetForm(initial={"status": Asset.Status.IN_USE, "condition": Asset.Condition.GOOD})
    return render(request, "assets/asset_form.html", {"form": form, "creating": True})


@it_required
def asset_edit(request, pk):
    asset = get_object_or_404(Asset, pk=pk)
    before = Asset.objects.get(pk=pk)
    if request.method == "POST":
        form = AssetForm(request.POST, instance=asset)
        if form.is_valid():
            try:
                services.update_asset(form.save(commit=False), before, request.user, form.cleaned_data.get("note", ""))
            except ValidationError as exc:
                _apply_errors(form, exc)
            else:
                messages.success(request, f"{asset.tag} updated.")
                return redirect(asset)
    else:
        form = AssetForm(instance=asset)
    return render(request, "assets/asset_form.html", {"form": form, "asset": asset, "creating": False})


@it_required
def asset_detail(request, pk):
    asset = get_object_or_404(Asset.objects.select_related("category", "department", "location", "assigned_to",
                                                           "supplier"), pk=pk)
    since = timezone.now() - timedelta(days=services.FAILURE_WINDOW_DAYS)
    fail = services.failure_counts(since, qs=Asset.objects.filter(pk=asset.pk)).first()
    failures = (fail.incidents + fail.repairs) if fail else 0
    asset.overdue_maintenance = asset.maintenance_records.filter(status="scheduled", due_date__lt=timezone.localdate()).count()
    return render(request, "assets/asset_detail.html", {
        "asset": asset, "history": services.service_history(asset), "expenses": services.expenses(asset),
        "plans": asset.maintenance_plans.select_related("responsible"),
        "upcoming": asset.maintenance_records.filter(status="scheduled").select_related("responsible"),
        "open_tickets": asset.tickets.open(), "attachments": asset.attachments.filter(is_deleted=False),
        "attach_form": AttachmentForm(), "reasons": services.attention_reasons(asset, failures),
        "ticket_count": asset.tickets.count(), "failures": failures,
        "scan_url": services.scan_url(request, asset),
    })


@it_required
def asset_status(request, pk):
    asset = get_object_or_404(Asset, pk=pk)
    if request.method == "POST":
        form = StatusForm(request.POST)
        if form.is_valid():
            try:
                services.change_status(asset, form.cleaned_data["status"], request.user, form.cleaned_data["reason"])
            except ValidationError as exc:
                _apply_errors(form, exc)
            else:
                messages.success(request, f"{asset.tag} is now {asset.get_status_display().lower()}.")
                return redirect(asset)
    else:
        form = StatusForm(initial={"status": request.GET.get("to") or Asset.Status.RETIRED})
    return render(request, "assets/asset_status.html", {"form": form, "asset": asset})


@it_required
@require_POST
def asset_attach(request, pk):
    asset = get_object_or_404(Asset, pk=pk)
    form = AttachmentForm(request.POST, request.FILES)
    if form.is_valid() and form.uploads:
        services.add_attachments(asset, form.uploads, request.user, form.cleaned_data["description"])
        messages.success(request, f"{len(form.uploads)} document(s) added.")
    else:
        for errs in form.errors.values():
            for e in errs:
                messages.error(request, e)
        if form.is_valid():
            messages.error(request, "Choose at least one file.")
    return redirect(reverse("assets:detail", args=[asset.pk]) + "#documents")


@it_required
def attachment_view(request, pk):
    att = get_object_or_404(AssetAttachment, pk=pk)
    if att.is_deleted:
        raise PermissionDenied("This file has been removed.")
    return serving.serve(request, att, download=request.GET.get("download") == "1")


@it_required
def attachment_hide(request, pk):
    from apps.tasks.forms import VoidForm

    att = get_object_or_404(AssetAttachment.objects.select_related("asset"), pk=pk)
    if request.method == "POST":
        form = VoidForm(request.POST)
        if form.is_valid():
            services.hide_attachment(att, request.user, form.cleaned_data["reason"])
            messages.success(request, "Document removed from view. It is kept for the audit record.")
            return redirect(att.asset)
    else:
        form = VoidForm()
    return render(request, "assets/simple_action.html", {"form": form, "asset": att.asset,
                                                         "heading": f"Remove “{att.original_name}”",
                                                         "button": "Remove from view", "button_class": "btn-danger"})


# ---------------------------------------------------------------------------
# QR codes
# ---------------------------------------------------------------------------
@it_required
def asset_qr(request, pk):
    asset = get_object_or_404(Asset, pk=pk)
    resp = HttpResponse(services.qr_png(services.scan_url(request, asset)), content_type="image/png")
    resp["Cache-Control"] = "private, max-age=3600"
    if request.GET.get("download"):
        resp["Content-Disposition"] = f'attachment; filename="{asset.tag}-qr.png"'
    return resp


@it_required
def asset_labels(request):
    """Printable sheet of QR labels for the selected (or filtered) assets."""
    ids = [int(i) for i in request.GET.getlist("id") if i.isdigit()]
    qs = Asset.objects.filter(pk__in=ids) if ids else _filtered(request)[0]
    assets = list(qs.select_related("location")[:200])
    import base64

    labels = [(a, base64.b64encode(services.qr_png(services.scan_url(request, a), box_size=5)).decode()) for a in assets]
    return render(request, "assets/labels.html", {"labels": labels})


@login_required
def asset_scan(request, token):
    """
    Target of the QR code. Signing in is required (login_required). IT staff
    see the asset record; everyone else is offered "report a problem with this
    device" — they never see the asset's details.
    """
    asset = get_object_or_404(Asset, qr_token=token)
    if is_it_staff(request.user):
        return redirect(asset)
    return redirect(reverse("helpdesk:create") + f"?asset={asset.qr_token}")


# ---------------------------------------------------------------------------
# Categories and suppliers
# ---------------------------------------------------------------------------
@role_required(Role.ADMIN)
def category_list(request):
    return render(request, "assets/category_list.html", {
        "items": AssetCategory.objects.annotate(n=Count("assets"))})


@role_required(Role.ADMIN)
def category_edit(request, pk=None):
    obj = get_object_or_404(AssetCategory, pk=pk) if pk else None
    if request.method == "POST":
        form = AssetCategoryForm(request.POST, instance=obj)
        if form.is_valid():
            item = form.save()
            audit.log(AuditLog.Action.CONFIG_CHANGE, item, message=f"Asset category {'updated' if obj else 'created'}")
            messages.success(request, f"Category '{item}' saved.")
            return redirect("assets:category_list")
    else:
        form = AssetCategoryForm(instance=obj)
    return render(request, "accounts/simple_form.html", {"form": form, "obj": obj, "label": "Asset category",
                                                         "label_plural": "Asset categories",
                                                         "list_url": "assets:category_list"})


@it_required
def supplier_list(request):
    return render(request, "assets/supplier_list.html", {"items": Supplier.objects.annotate(n=Count("assets"))})


@it_required
def supplier_edit(request, pk=None):
    obj = get_object_or_404(Supplier, pk=pk) if pk else None
    if request.method == "POST":
        form = SupplierForm(request.POST, instance=obj)
        if form.is_valid():
            item = form.save()
            audit.log(AuditLog.Action.UPDATE if obj else AuditLog.Action.CREATE, item, message="Supplier saved")
            messages.success(request, f"Supplier '{item}' saved.")
            return redirect("assets:supplier_list")
    else:
        form = SupplierForm(instance=obj)
    return render(request, "accounts/simple_form.html", {"form": form, "obj": obj, "label": "Supplier",
                                                         "list_url": "assets:supplier_list"})
