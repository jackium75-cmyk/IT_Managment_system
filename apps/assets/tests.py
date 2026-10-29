import io
import shutil
import tempfile
from datetime import timedelta
from decimal import Decimal

from django.core.exceptions import ValidationError
from django.test import override_settings
from django.urls import reverse
from django.utils import timezone

from apps.accounts.roles import Role
from apps.core.models import AuditLog, SystemSetting
from apps.core.testing import HotelTestCase
from apps.diary.models import ActivityLog
from apps.helpdesk import services as hd
from apps.helpdesk.models import Ticket, TicketCategory
from apps.maintenance import services as maint
from apps.maintenance.models import MaintenancePlan, MaintenanceRecord
from apps.notifications.models import Notification
from apps.tasks.models import Task

from . import services
from .models import Asset, AssetCategory

TMP_MEDIA = tempfile.mkdtemp(prefix="hotel-it-test-assets-")
R = MaintenanceRecord.Status


@override_settings(MEDIA_ROOT=TMP_MEDIA)
class AssetTestCase(HotelTestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(TMP_MEDIA, ignore_errors=True)

    def setUp(self):
        super().setUp()
        self.it = self.make_user("ian", Role.IT_OFFICER)
        self.admin = self.make_user("adm", Role.ADMIN)
        self.emp = self.make_user("rose")
        self.mgr = self.make_user("fatma", Role.MANAGER)
        self.printers = AssetCategory.objects.get(code="PRN")
        self.today = timezone.localdate()

    def asset(self, **kw):
        kw.setdefault("name", "Reception printer")
        kw.setdefault("category", self.printers)
        kw.setdefault("department", self.dept)
        return services.create_asset(Asset(**kw), self.it)

    def plan(self, asset, **kw):
        kw.setdefault("title", "Clean and inspect")
        kw.setdefault("interval_value", 3)
        kw.setdefault("interval_unit", "months")
        kw.setdefault("next_due_date", self.today + timedelta(days=20))
        kw.setdefault("lead_days", 7)
        kw.setdefault("responsible", self.it)
        kw.setdefault("checklist", "Clean rollers\nCheck toner\nPrint test page")
        return maint.create_plan(MaintenancePlan(asset=asset, **kw), self.it)


class AssetRegisterTests(AssetTestCase):
    def test_tags_are_sequential_per_category(self):
        a, b = self.asset(), self.asset(name="Back office printer")
        pc = self.asset(name="Reception PC", category=AssetCategory.objects.get(code="PC"))
        self.assertEqual((a.tag, b.tag, pc.tag), ("PRN-0001", "PRN-0002", "PC-0001"))
        self.assertTrue(AuditLog.objects.filter(action=AuditLog.Action.CREATE, object_id=str(a.pk)).exists())

    def test_create_via_form_and_validation(self):
        self.login(self.it)
        resp = self.client.post(reverse("assets:create"), {
            "name": "Core switch", "category": AssetCategory.objects.get(code="SW").pk, "status": "in_use",
            "condition": "good", "ip_address": "10.0.0.2", "mac_address": "AA:BB:CC:00:11:22",
            "purchase_cost": "1200.50", "replacement": "none"})
        asset = Asset.objects.get(name="Core switch")
        self.assertRedirects(resp, asset.get_absolute_url())
        self.assertEqual(asset.currency, SystemSetting.get().currency)  # defaults to hotel currency
        resp = self.client.post(reverse("assets:create"), {
            "name": "Bad", "category": self.printers.pk, "status": "in_use", "condition": "good",
            "mac_address": "not-a-mac", "ip_address": "999.1.1.1", "replacement": "none"})
        self.assertEqual(resp.status_code, 200)
        self.assertIn("mac_address", resp.context["form"].errors)
        self.assertIn("ip_address", resp.context["form"].errors)

    def test_credentials_are_refused(self):
        for text in ("admin password: Hotel123", "WiFi PSK=supersecret", "licence key: ABCD-1234", "pwd=letmein"):
            with self.assertRaises(ValidationError, msg=text):
                self.asset(notes=text)
        self.asset(notes="Password is kept in the IT password manager")  # mentioning passwords is fine

    def test_duplicate_serial_rejected(self):
        self.asset(serial_number="SN123", manufacturer="HP")
        self.login(self.it)
        resp = self.client.post(reverse("assets:create"), {
            "name": "Dup", "category": self.printers.pk, "status": "in_use", "condition": "good",
            "serial_number": "sn123", "manufacturer": "HP", "replacement": "none"})
        self.assertIn("serial_number", resp.context["form"].errors)

    def test_moves_create_assignment_history(self):
        a = self.asset()
        before = Asset.objects.get(pk=a.pk)
        a.assigned_to, a.department = self.emp, self.other_dept
        services.update_asset(a, before, self.it, note="Moved to F&B office")
        h = a.assignment_history.first()
        self.assertEqual((h.to_user, h.to_department, h.from_department, h.note),
                         (self.emp, self.other_dept, self.dept, "Moved to F&B office"))
        self.assertTrue(ActivityLog.objects.filter(asset=a, event=ActivityLog.Event.ASSET_UPDATED).exists())

    def test_retire_requires_reason_and_stops_maintenance(self):
        a = self.asset()
        self.plan(a)
        before = Asset.objects.get(pk=a.pk)
        a.status = Asset.Status.RETIRED
        with self.assertRaises(ValidationError):
            services.update_asset(a, before, self.it)  # must use change_status
        with self.assertRaises(ValidationError):
            services.change_status(before, Asset.Status.RETIRED, self.it, "")
        services.change_status(before, Asset.Status.RETIRED, self.it, "Beyond economical repair")
        self.assertFalse(MaintenancePlan.objects.get(asset=a).is_active)
        self.assertEqual(MaintenanceRecord.objects.get(asset=a).status, R.CANCELLED)

    def test_only_it_staff_can_open_the_register(self):
        a = self.asset()
        for user in (self.emp, self.mgr):
            self.login(user)
            for url in (reverse("assets:list"), a.get_absolute_url(), reverse("assets:qr", args=[a.pk]),
                        reverse("assets:export", args=["csv"]), reverse("maintenance:list")):
                self.assertEqual(self.client.get(url).status_code, 403, (user.username, url))

    def test_export_csv_and_xlsx(self):
        self.asset(name="=HYPERLINK(\"http://evil\")")
        self.login(self.it)
        resp = self.client.get(reverse("assets:export", args=["csv"]))
        body = resp.content.decode("utf-8-sig")
        self.assertIn("Asset ID,Name", body)
        self.assertIn("'=HYPERLINK", body)  # formula injection neutralised
        resp = self.client.get(reverse("assets:export", args=["xlsx"]))
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(resp.content)).active
        self.assertEqual(ws["A2"].value, "PRN-0001")
        self.assertTrue(AuditLog.objects.filter(action=AuditLog.Action.REPORT).exists())


class QRCodeTests(AssetTestCase):
    def test_qr_scan_routes_by_role(self):
        a = self.asset()
        url = reverse("assets:scan", args=[a.qr_token])
        self.assertIn(str(a.qr_token), url)
        self.assertNotIn(a.tag, url)  # the QR code reveals nothing about the asset
        self.assertEqual(self.client.get(url).status_code, 302)  # sign-in required
        self.assertIn(reverse("accounts:login"), self.client.get(url).url)
        self.login(self.it)
        self.assertRedirects(self.client.get(url), a.get_absolute_url())
        self.login(self.emp)
        resp = self.client.get(url)
        self.assertRedirects(resp, reverse("helpdesk:create") + f"?asset={a.qr_token}")
        page = self.client.get(resp.url)
        self.assertContains(page, a.tag)
        self.assertNotContains(page, "Serial number")

    def test_employee_reports_problem_from_qr(self):
        a = self.asset()
        self.login(self.emp)
        cat = TicketCategory.objects.get(name="Printer / scanner", parent=None)
        self.client.post(reverse("helpdesk:create"), {
            "asset_token": str(a.qr_token), "ticket_type": "incident", "subject": "Paper jam", "category": cat.pk,
            "description": "Jammed", "impact": "one", "urgency": "slowed"})
        ticket = Ticket.objects.get(subject="Paper jam")
        self.assertEqual(ticket.asset, a)
        self.assertIn(a.tag, ticket.affected_device)
        self.assertEqual(services.service_history(a)[0]["kind"], "incident")

    def test_qr_png_and_labels(self):
        a = self.asset()
        self.login(self.it)
        resp = self.client.get(reverse("assets:qr", args=[a.pk]))
        self.assertEqual(resp["Content-Type"], "image/png")
        self.assertTrue(resp.content.startswith(b"\x89PNG"))
        self.assertContains(self.client.get(reverse("assets:labels") + f"?id={a.pk}"), a.tag)


class MaintenanceTests(AssetTestCase):
    def test_plan_creates_one_scheduled_visit_and_task_within_lead_time(self):
        a = self.asset()
        plan = self.plan(a, next_due_date=self.today + timedelta(days=20))
        record = MaintenanceRecord.objects.get(plan=plan)
        self.assertEqual(record.status, R.SCHEDULED)
        self.assertIsNone(record.task)  # 20 days away, lead time 7
        maint.run_jobs(self.today + timedelta(days=14))
        record.refresh_from_db()
        self.assertIsNotNone(record.task)
        self.assertEqual((record.task.asset, record.task.assigned_to), (a, self.it))
        self.assertTrue(Notification.objects.filter(recipient=self.it, kind=Notification.Kind.MAINTENANCE).exists())
        maint.run_jobs(self.today + timedelta(days=15))
        self.assertEqual(Task.objects.filter(asset=a).count(), 1)  # never duplicated
        self.assertEqual(MaintenanceRecord.objects.filter(plan=plan, status=R.SCHEDULED).count(), 1)

    def test_completion_records_checklist_and_schedules_next(self):
        a = self.asset()
        plan = self.plan(a, next_due_date=self.today)
        record = MaintenanceRecord.objects.get(plan=plan)
        self.assertIsNotNone(record.task)  # due today: task created immediately
        with self.assertRaises(ValidationError):
            maint.complete_record(record, self.it, completed_on=self.today, work_done="Cleaned",
                                  checklist_results=[{"item": "Check toner", "result": "issue", "note": ""}])
        maint.complete_record(record, self.it, completed_on=self.today, work_done="Cleaned rollers",
                              checklist_results=[{"item": "Clean rollers", "result": "ok", "note": ""},
                                                 {"item": "Check toner", "result": "issue", "note": "Low toner"}],
                              cost=Decimal("15.00"), asset_condition="fair")
        record.refresh_from_db()
        self.assertEqual((record.status, record.on_time, len(record.issues)), (R.COMPLETED, True, 1))
        self.assertEqual(record.task.status, Task.Status.COMPLETED)
        a.refresh_from_db()
        self.assertEqual(a.condition, "fair")
        plan.refresh_from_db()
        self.assertEqual(plan.next_due_date, plan.following(self.today))
        nxt = MaintenanceRecord.objects.get(plan=plan, status=R.SCHEDULED)
        self.assertEqual(nxt.due_date, plan.next_due_date)
        self.assertEqual(services.expenses(a), {"USD": Decimal("15.00")})
        self.assertTrue(ActivityLog.objects.filter(asset=a, event=ActivityLog.Event.MAINTENANCE_DONE).exists())

    def test_complete_via_form(self):
        a = self.asset()
        plan = self.plan(a, next_due_date=self.today)
        record = plan.records.get()
        self.login(self.it)
        resp = self.client.post(reverse("maintenance:complete", args=[record.pk]), {
            "completed_on": self.today.isoformat(), "work_done": "Done", "check_0": "ok", "check_1": "issue",
            "note_1": "", "check_2": "na"})
        self.assertEqual(resp.status_code, 200)
        self.assertIn("note_1", resp.context["form"].errors)
        resp = self.client.post(reverse("maintenance:complete", args=[record.pk]), {
            "completed_on": self.today.isoformat(), "work_done": "Done", "check_0": "ok", "check_1": "issue",
            "note_1": "Toner at 5%", "check_2": "na"})
        self.assertRedirects(resp, record.get_absolute_url())

    def test_missed_after_grace_and_next_visit_scheduled(self):
        a = self.asset()
        plan = self.plan(a, next_due_date=self.today - timedelta(days=1))
        grace = SystemSetting.get().maintenance_grace_days
        maint.run_jobs(self.today)
        record = plan.records.get(due_date=self.today - timedelta(days=1))
        self.assertEqual(record.status, R.SCHEDULED)
        self.assertTrue(Notification.objects.filter(kind=Notification.Kind.OVERDUE, recipient=self.it).exists())
        maint.run_jobs(self.today + timedelta(days=grace))
        record.refresh_from_db()
        self.assertEqual(record.status, R.MISSED)
        self.assertEqual(record.task.status, Task.Status.CANCELLED)
        nxt = plan.records.get(status=R.SCHEDULED)
        self.assertEqual(nxt.due_date, plan.following(self.today - timedelta(days=1)))

    def test_reschedule_keeps_original_date_for_compliance(self):
        a = self.asset()
        plan = self.plan(a, next_due_date=self.today - timedelta(days=3))
        record = plan.records.get()
        with self.assertRaises(ValidationError):
            maint.reschedule_record(record, self.it, self.today + timedelta(days=2), "")
        maint.reschedule_record(record, self.it, self.today + timedelta(days=2), "Server room locked for audit")
        record.refresh_from_db()
        self.assertEqual((record.original_due_date, record.reschedule_count), (self.today - timedelta(days=3), 1))
        maint.complete_record(record, self.it, completed_on=self.today, work_done="Done")
        record.refresh_from_db()
        self.assertFalse(record.on_time)  # completed after the ORIGINAL due date
        c = maint.compliance(self.today - timedelta(days=10), self.today)
        self.assertEqual((c["due"], c["late"], c["on_time"], c["rescheduled"], c["rate"]), (1, 1, 0, 1, 0))

    def test_corrective_repair_from_ticket(self):
        a = self.asset()
        ticket = hd.create_ticket(Ticket(subject="Printer dead", description="x",
                                         category=TicketCategory.objects.get(name="Printer / scanner", parent=None),
                                         asset=a), self.emp)
        rec = maint.record_corrective(a, self.it, title="Replaced fuser", work_done="New fuser fitted",
                                      completed_on=self.today, cost=Decimal("80"), ticket=ticket,
                                      asset_status="in_use")
        self.assertEqual((rec.kind, rec.status, rec.ticket), ("corrective", R.COMPLETED, ticket))
        since = timezone.now() - timedelta(days=30)
        row = services.failure_counts(since).get(pk=a.pk)
        self.assertEqual((row.incidents, row.repairs), (1, 1))
        kinds = {e["kind"] for e in services.service_history(a)}
        self.assertTrue({"incident", "repair"} <= kinds)

    def test_attention_rules(self):
        a = self.asset(condition="poor", warranty_expiry=self.today + timedelta(days=10))
        b = self.asset(name="Fine printer")
        self.plan(b, next_due_date=self.today - timedelta(days=2))
        found = dict((x.pk, r) for x, r in services.needing_attention())
        self.assertIn("condition: poor", found[a.pk])
        self.assertTrue(any("warranty" in r for r in found[a.pk]))
        self.assertIn("maintenance overdue", found[b.pk])

    def test_pages_render(self):
        a = self.asset()
        plan = self.plan(a, next_due_date=self.today)
        record = plan.records.get()
        self.login(self.it)
        for url in (reverse("assets:list"), reverse("assets:list") + "?attention=1&status=all", reverse("assets:create"),
                    a.get_absolute_url(), reverse("assets:edit", args=[a.pk]), reverse("assets:status", args=[a.pk]),
                    reverse("assets:supplier_list"), reverse("maintenance:list"), reverse("maintenance:list") + "?view=overdue",
                    reverse("maintenance:calendar"), reverse("maintenance:compliance"), reverse("maintenance:plan_create"),
                    reverse("maintenance:plan_edit", args=[plan.pk]), record.get_absolute_url(),
                    reverse("maintenance:complete", args=[record.pk]), reverse("maintenance:reschedule", args=[record.pk]),
                    reverse("maintenance:corrective_create") + f"?asset={a.pk}", reverse("dashboard:index"),
                    reverse("tasks:create") + f"?asset={a.pk}", reverse("diary:list") + f"?asset={a.tag}"):
            self.assertEqual(self.client.get(url).status_code, 200, url)
        self.login(self.admin)
        self.assertEqual(self.client.get(reverse("assets:category_list")).status_code, 200)
